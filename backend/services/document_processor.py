from pathlib import Path
from typing import List, Dict, Any
import pymupdf
from docx import Document


def extract_text_from_pdf(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text per page from a PDF file.
    Preserves page number metadata and handles empty or scanned pages.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        doc = pymupdf.open(str(path))
    except Exception as e:
        raise ValueError(f"Failed to open PDF document: {str(e)}")

    pages: List[Dict[str, Any]] = []

    try:
        for page_number, page in enumerate(doc, start=1):
            text = page.get_text()
            # Normalize whitespace while preserving structure
            cleaned_text = text.strip() if text else ""
            pages.append({
                "page_number": page_number,
                "text": cleaned_text
            })
    finally:
        doc.close()

    return pages


def extract_text_from_docx(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text from a DOCX document including paragraphs and tables.
    Preserves paragraph breaks for chunking.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        doc = Document(str(path))
    except Exception as e:
        raise ValueError(f"Failed to open DOCX document: {str(e)}")

    text_parts: List[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_parts.append(text)

    # Also capture table contents if present
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                text_parts.append(" | ".join(row_cells))

    full_text = "\n\n".join(text_parts)

    return [
        {
            "page_number": 1,
            "text": full_text
        }
    ]


def extract_text_from_txt(file_path: str) -> List[Dict[str, Any]]:
    """
    Extracts text from a plain text (.txt) or markdown (.md) file.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
    except Exception as e:
        raise ValueError(f"Failed to read text file: {str(e)}")

    return [
        {
            "page_number": 1,
            "text": text.strip()
        }
    ]